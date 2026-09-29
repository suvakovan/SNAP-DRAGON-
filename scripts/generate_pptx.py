"""
Generates a 12-Slide PowerPoint (.pptx) Presentation for LectureLens.
Styled in iQOO Cyber Yellow (#FFD100) & Obsidian Black (#0A0A0C).
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

root_dir = Path(__file__).resolve().parent.parent

# Theme Colors
COLOR_BG = RGBColor(10, 10, 12)         # Obsidian Black #0A0A0C
COLOR_CARD = RGBColor(24, 24, 28)       # Dark Charcoal #18181C
COLOR_YELLOW = RGBColor(255, 209, 0)     # iQOO Cyber Yellow #FFD100
COLOR_WHITE = RGBColor(255, 255, 255)   # Crisp White #FFFFFF
COLOR_GRAY = RGBColor(161, 161, 170)    # Subtitle Gray #A1A1AA

def apply_slide_bg(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

def add_header(slide, title_text, category_text="LECTURELENS — QUALCOMM SNAPDRAGON AI CHALLENGE"):
    # Category / Tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_YELLOW

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ----------------------------------------------------
    # SLIDE 1: Title Slide
    # ----------------------------------------------------
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide1)

    # Big Yellow Accent Card
    accent_card = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(6.0), Inches(5.0))
    accent_card.fill.solid()
    accent_card.fill.fore_color.rgb = COLOR_CARD
    accent_card.line.color.rgb = COLOR_YELLOW
    accent_card.line.width = Pt(2)

    tb1 = slide1.shapes.add_textbox(Inches(1.1), Inches(1.5), Inches(5.4), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "LECTURELENS"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLOR_YELLOW

    p2 = tf1.add_paragraph()
    p2.text = "Fully Offline Lecture Copilot Accelerated for Snapdragon PCs"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.space_before = Pt(14)

    p3 = tf1.add_paragraph()
    p3.text = "Qualcomm Snapdragon AI Lab Challenge Submission\n100% Private • NPU Accelerated • Zero Cloud Dependency"
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_GRAY
    p3.space_before = Pt(20)

    # Preview Image on Slide 1
    preview_img = root_dir / "docs" / "assets" / "app_preview.png"
    if preview_img.exists():
        slide1.shapes.add_picture(str(preview_img), Inches(7.1), Inches(1.8), width=Inches(5.4))

    # ----------------------------------------------------
    # SLIDE 2: The Problem
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide2)
    add_header(slide2, "The Problem: Classroom Wi-Fi Deadzones & Cloud Latency")

    cards_data_s2 = [
        ("🌐 Campus Wi-Fi Deadzones", "Overcrowded lecture halls suffer from severe bandwidth throttling or zero cellular coverage, rendering cloud AI tools unusable."),
        ("⚡ Fast-Paced STEM Lectures", "Engineering & STEM students struggle to listen attentively while scribbling complex equations and technical terminology."),
        ("🔒 Privacy & IP Concerns", "Uploading proprietary course lectures or confidential research audio to public cloud APIs violates data privacy policies.")
    ]

    for idx, (head, desc) in enumerate(cards_data_s2):
        left = Inches(0.8 + idx * 3.9)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.6), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_YELLOW if idx == 0 else COLOR_CARD
        card.line.width = Pt(1.5)

        tb = slide2.shapes.add_textbox(left + Inches(0.2), Inches(2.2), Inches(3.2), Inches(4.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = head
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_YELLOW

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(12)

    # ----------------------------------------------------
    # SLIDE 3: The Solution
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide3)
    add_header(slide3, "The Solution: LectureLens On-Device AI Copilot")

    pillars = [
        ("🔒 100% Data Privacy", "All audio processing, neural transcription, and LLM summary generation happen locally on the laptop with zero cloud requests."),
        ("⚡ Hexagon NPU Acceleration", "Optimized with Qualcomm QNN Execution Provider (HTP) & Foundry Local NPU model endpoints for maximum battery efficiency."),
        ("🎓 Automated Study Artifacts", "Instant conversion of raw classroom speech into executive notes, key term glossaries, interactive quizzes, and flashcards.")
    ]

    for idx, (title, text) in enumerate(pillars):
        top = Inches(1.8 + idx * 1.7)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_YELLOW
        card.line.width = Pt(1)

        tb = slide3.shapes.add_textbox(Inches(1.1), top + Inches(0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_YELLOW

        p2 = tf.add_paragraph()
        p2.text = text
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 4: Key Capabilities & Student Workflow
    # ----------------------------------------------------
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide4)
    add_header(slide4, "Student Workflow: From Audio Stream to Knowledge Mastery")

    steps = [
        ("1. Record / Upload", "Live mic capture or WAV/MP3/FLAC file import"),
        ("2. Transcribe", "Neural Whisper STT with sub-second chunking"),
        ("3. Summarize", "Executive bullets, key terms & revision paragraph"),
        ("4. Quiz & Flashcards", "Interactive multiple choice Q&A & active recall cards"),
        ("5. Hybrid Search", "Cross-lecture vector embeddings & BM25 keyword index")
    ]

    for idx, (stitle, sdesc) in enumerate(steps):
        left = Inches(0.8 + idx * 2.38)
        box = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(2.2), Inches(2.2), Inches(4.0))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD
        box.line.color.rgb = COLOR_YELLOW
        box.line.width = Pt(1.5)

        tb = slide4.shapes.add_textbox(left + Inches(0.1), Inches(2.3), Inches(2.0), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = stitle
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_YELLOW

        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(10)

    # ----------------------------------------------------
    # SLIDE 5: System Architecture & Data Pipeline
    # ----------------------------------------------------
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide5)
    add_header(slide5, "System Architecture & On-Device Data Pipeline")

    arch_blocks = [
        ("Audio Resampler & VAD", "16 kHz mono conversion with energy-based voice activity detection."),
        ("Whisper STT (QNN/HTP)", "ONNX Runtime QNN provider for NPU speech recognition + PyTorch CPU fallback."),
        ("Foundry Local LLM", "Standardized local NPU chat endpoint + Ollama local CPU fallback."),
        ("Hybrid Search Engine", "ONNX MiniLM passage embeddings + BM25 keyword index.")
    ]

    for idx, (bhead, bdesc) in enumerate(arch_blocks):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(2.0 + row * 2.4)

        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.7), Inches(2.0))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_YELLOW
        card.line.width = Pt(1)

        tb = slide5.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(5.3), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = bhead
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_YELLOW

        p2 = tf.add_paragraph()
        p2.text = bdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(8)

    # ----------------------------------------------------
    # SLIDE 6: Snapdragon NPU Acceleration Design
    # ----------------------------------------------------
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide6)
    add_header(slide6, "Snapdragon NPU Acceleration & Dual Engine Architecture")

    npu_info = [
        ("⚡ Qualcomm Hexagon NPU Offloading", "Offloads dense matrix multiplications to Hexagon NPU to extend battery life and keep CPU thermals cool during long 2-hour lectures."),
        ("🛠️ ONNX Runtime + QNN Provider", "`stt_qnn.py` connects directly to Qualcomm HTP driver for NPU-accelerated Whisper speech recognition."),
        ("💬 Foundry Local NPU LLM Endpoint", "`llm_foundry.py` wraps local `/v1/chat/completions` API targeting Snapdragon NPU LLM models."),
        ("💻 Graceful CPU Fallback Architecture", "Zero-downtime fallback to PyTorch & Ollama CPU engines when QNN hardware is absent on non-Snapdragon dev machines.")
    ]

    for idx, (title, desc) in enumerate(npu_info):
        top = Inches(1.8 + idx * 1.3)
        card = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_YELLOW
        card.line.width = Pt(1)

        tb = slide6.shapes.add_textbox(Inches(1.0), top + Inches(0.1), Inches(11.3), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_YELLOW

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(2)

    # ----------------------------------------------------
    # SLIDE 7: Verified Performance & Benchmark Metrics
    # ----------------------------------------------------
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide7)
    add_header(slide7, "Honest & Verified Performance Benchmarks")

    metrics_data = [
        ("10% WER", "STT Accuracy", "Verified against 40-word human TTS reference transcript"),
        ("4.93 tok/sec", "LLM Speed (CPU)", "Measured on Ollama phi3:mini local CPU engine"),
        ("~0.166 RTF", "STT Processing", "Real-time factor measured for Whisper CPU inference"),
        ("24 / 24", "Unit & Integration Tests", "100% test suite pass rate verified with pytest")
    ]

    for idx, (val, label, sub) in enumerate(metrics_data):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.9 + row * 2.5)

        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_YELLOW
        card.line.width = Pt(1.5)

        tb = slide7.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(5.3), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = val
        p1.font.size = Pt(36)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_YELLOW

        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_WHITE

        p3 = tf.add_paragraph()
        p3.text = sub
        p3.font.size = Pt(12)
        p3.font.color.rgb = COLOR_GRAY
        p3.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 8: Privacy & Strict Offline Guarantee
    # ----------------------------------------------------
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide8)
    add_header(slide8, "Strict Offline Guarantee & Zero Cloud Interception")

    priv_cards = [
        ("0 Cloud Network Calls", "Confirmed by `scripts/offline_check.py` socket guard monkeypatching."),
        ("100% Data Residency", "Audio, transcripts, notes, and vector databases remain encrypted on local disk."),
        ("Airplane Mode Ready", "Operates reliably without active Wi-Fi, cellular, or Bluetooth networks.")
    ]

    for idx, (head, body) in enumerate(priv_cards):
        left = Inches(0.8 + idx * 3.9)
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.6), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_YELLOW
        card.line.width = Pt(1.5)

        tb = slide8.shapes.add_textbox(left + Inches(0.2), Inches(2.2), Inches(3.2), Inches(4.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = head
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_YELLOW

        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.size = Pt(14)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(14)

    # ----------------------------------------------------
    # SLIDE 9: User Interface & iQOO Cyber Yellow Theme
    # ----------------------------------------------------
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide9)
    add_header(slide9, "User Experience: iQOO Cyber Yellow & Obsidian Black UI")

    tb9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.5), Inches(5.0))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    p = tf9.paragraphs[0]
    p.text = "🎨 High-Contrast Streamlit Web Dashboard"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_YELLOW

    bullets = [
        "Inspired by flagship iQOO Monster Cyber Yellow (#FFD100) & Pitch Carbon (#000000) styling.",
        "8 Responsive Navigation Tabs: Record, Notes, Quiz, Flashcards, Search, Q&A, Performance, History.",
        "Real-Time Acceleration Sidebar: Shows active hardware status (QNN NPU vs CPU Fallback).",
        "Accessible & High-Contrast: Optimized typography for fast classroom readability."
    ]

    for b in bullets:
        bp = tf9.add_paragraph()
        bp.text = "• " + b
        bp.font.size = Pt(13)
        bp.font.color.rgb = COLOR_WHITE
        bp.space_before = Pt(10)

    if preview_img.exists():
        slide9.shapes.add_picture(str(preview_img), Inches(6.6), Inches(1.8), width=Inches(5.9))

    # ----------------------------------------------------
    # SLIDE 10: Real World Impact & Target Audience
    # ----------------------------------------------------
    slide10 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide10)
    add_header(slide10, "Target Persona & Real-World Educational Impact")

    aud_data = [
        ("🎓 STEM Students", "Rapid extraction of complex engineering concepts, formulas, and technical glossaries."),
        ("⚖️ Law & Medical Students", "High-density active recall quiz generation and automatic flashcard creation."),
        ("🔬 Research Faculty", "Secure, confidential lecture recording and seminar analysis with zero privacy leaks.")
    ]

    for idx, (atitle, adesc) in enumerate(aud_data):
        top = Inches(1.8 + idx * 1.7)
        card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_YELLOW
        card.line.width = Pt(1)

        tb = slide10.shapes.add_textbox(Inches(1.1), top + Inches(0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = atitle
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_YELLOW

        p2 = tf.add_paragraph()
        p2.text = adesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 11: Future Roadmap
    # ----------------------------------------------------
    slide11 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide11)
    add_header(slide11, "Future Roadmap: Extending On-Device Intelligence")

    roadmap_items = [
        ("1. Speaker Diarization", "On-device voice embeddings to separate professor vs student questions in live audio."),
        ("2. Multilingual Support", "Real-time speech translation for Tamil, Hindi, Spanish, and German lectures."),
        ("3. Wi-Fi Direct Sync", "Local peer-to-peer note sharing between Snapdragon laptops and tablets without internet.")
    ]

    for idx, (rtitle, rdesc) in enumerate(roadmap_items):
        top = Inches(1.8 + idx * 1.7)
        card = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_YELLOW
        card.line.width = Pt(1)

        tb = slide11.shapes.add_textbox(Inches(1.1), top + Inches(0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = rtitle
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_YELLOW

        p2 = tf.add_paragraph()
        p2.text = rdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 12: Conclusion & Q&A
    # ----------------------------------------------------
    slide12 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_slide_bg(slide12)

    card12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(11.7), Inches(5.5))
    card12.fill.solid()
    card12.fill.fore_color.rgb = COLOR_CARD
    card12.line.color.rgb = COLOR_YELLOW
    card12.line.width = Pt(2)

    tb12 = slide12.shapes.add_textbox(Inches(1.2), Inches(1.4), Inches(10.9), Inches(4.7))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p1 = tf12.paragraphs[0]
    p1.text = "LECTURELENS"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_YELLOW
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf12.add_paragraph()
    p2.text = "Bringing Private, Offline AI to Every Engineering Classroom."
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(16)

    p3 = tf12.add_paragraph()
    p3.text = "GitHub Repo: https://github.com/suvakovan/SNAP-DRAGON-\nQualcomm Snapdragon AI Lab Challenge 2026"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_GRAY
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(24)

    p4 = tf12.add_paragraph()
    p4.text = "THANK YOU! QUESTIONS & LIVE DEMO"
    p4.font.size = Pt(22)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_YELLOW
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(30)

    output_path = root_dir / "LectureLens_Presentation.pptx"
    prs.save(str(output_path))
    print(f"✅ Successfully created 12-slide presentation at: {output_path}")

if __name__ == "__main__":
    create_deck()
